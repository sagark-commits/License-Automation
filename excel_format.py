"""OpenPyXL styling and visualizations for TMONE utilization workbooks."""

from __future__ import annotations

import calendar
from collections import defaultdict
from datetime import date, datetime
from typing import Any

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
SUBHEADER_FILL = PatternFill("solid", fgColor="2E75B6")
SUBHEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=9)
BANNER_FILL = PatternFill("solid", fgColor="1F4E79")
BANNER_FONT = Font(name="Calibri", bold=True, size=16, color="FFFFFF")
SUBTITLE_FONT = Font(name="Calibri", size=10, color="D6E4F0")
TITLE_FONT = Font(name="Calibri", bold=True, size=14, color="1F4E79")
LABEL_FONT = Font(name="Calibri", bold=True, size=10)
KPI_LABEL_FONT = Font(name="Calibri", bold=True, size=9, color="5B5B5B")
KPI_VALUE_FONT = Font(name="Calibri", bold=True, size=18, color="1F4E79")
SUMMARY_FILL = PatternFill("solid", fgColor="FFF2CC")
ALT_FILL = PatternFill("solid", fgColor="F5F7FA")
PEAK_FILL = PatternFill("solid", fgColor="C6EFCE")
WEEKEND_FILL = PatternFill("solid", fgColor="E8EEF4")
WEEKEND_HEADER = PatternFill("solid", fgColor="5B7FA6")
AGENT_ROW_FILL = PatternFill("solid", fgColor="DDEBF7")
SUP_ROW_FILL = PatternFill("solid", fgColor="FCE4D6")
WB_ROW_FILL = PatternFill("solid", fgColor="E4DFEC")
KPI_FILL = PatternFill("solid", fgColor="F8FAFC")
ZERO_FONT = Font(name="Calibri", color="BFBFBF", size=9)
DATA_FONT = Font(name="Calibri", size=10)
THIN = Side(style="thin", color="D9D9D9")
MEDIUM = Side(style="medium", color="1F4E79")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
KPI_BORDER = Border(left=THIN, right=THIN, top=MEDIUM, bottom=THIN)

LICENSE_ROW_FILLS = {
    "Ameyo Express": AGENT_ROW_FILL,
    "Ameyo Supervisor": SUP_ROW_FILL,
    "Ameyo Wallboard": WB_ROW_FILL,
}

HOUR_COLUMNS = [f"Count at {h} hour" for h in range(1, 25)]


def _month_label(month: str) -> str:
    y, m = map(int, month.split("-"))
    return f"{calendar.month_name[m]} {y}"


def _is_blank_row(ws, row: int, max_col: int) -> bool:
    return all(ws.cell(row, c).value in (None, "") for c in range(1, max_col + 1))


def _set(ws, row: int, col: int, **kwargs) -> None:
    cell = ws.cell(row, col)
    for key, value in kwargs.items():
        setattr(cell, key, value)


def _license_row_fill(license_label: str | None) -> PatternFill | None:
    if not license_label:
        return None
    return LICENSE_ROW_FILLS.get(str(license_label).strip())


def _weekend_columns(ws, first_date_col: int, max_col: int) -> set[int]:
    weekend_cols: set[int] = set()
    for col in range(first_date_col, max_col + 1):
        day_name = ws.cell(3, col).value
        if day_name in ("Sat", "Sun"):
            weekend_cols.add(col)
    return weekend_cols


def format_summary_sheet(ws, month: str) -> None:
    max_col = ws.max_column
    max_row = ws.max_row
    if max_col < 4:
        return

    title = f"TmOne License Utilization — {_month_label(month)}"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    _set(
        ws,
        1,
        1,
        value=title,
        font=BANNER_FONT,
        fill=BANNER_FILL,
        alignment=Alignment(horizontal="center", vertical="center"),
    )
    ws.row_dimensions[1].height = 34

    for col in range(1, 4):
        _set(
            ws,
            2,
            col,
            fill=HEADER_FILL,
            font=HEADER_FONT,
            alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
            border=BORDER,
        )

    weekend_cols = _weekend_columns(ws, 4, max_col)
    for col in range(4, max_col + 1):
        hdr_fill = WEEKEND_HEADER if col in weekend_cols else HEADER_FILL
        sub_fill = WEEKEND_FILL if col in weekend_cols else SUBHEADER_FILL
        _set(
            ws,
            2,
            col,
            fill=hdr_fill,
            font=HEADER_FONT,
            alignment=Alignment(horizontal="center", vertical="center"),
            border=BORDER,
        )
        _set(
            ws,
            3,
            col,
            fill=sub_fill,
            font=SUBHEADER_FONT,
            alignment=Alignment(horizontal="center", vertical="center"),
            border=BORDER,
        )

    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 18
    ws.freeze_panes = "D4"

    legend_row = max_row + 2
    _set(ws, legend_row, 1, value="Legend:", font=LABEL_FONT)
    for idx, (label, fill) in enumerate(LICENSE_ROW_FILLS.items(), start=2):
        _set(ws, legend_row, idx, value=label, fill=fill, font=DATA_FONT, border=BORDER,
             alignment=Alignment(horizontal="center"))
    _set(ws, legend_row, 5, value="Peak day", fill=PEAK_FILL, font=DATA_FONT, border=BORDER,
         alignment=Alignment(horizontal="center"))

    group_idx = 0
    for row in range(4, max_row + 1):
        if _is_blank_row(ws, row, max_col):
            ws.row_dimensions[row].height = 6
            continue

        group_idx += 1
        license_label = ws.cell(row, 2).value
        license_fill = _license_row_fill(license_label)
        row_fill = license_fill or (ALT_FILL if group_idx % 2 == 0 else None)

        project_val = ws.cell(row, 1).value
        if project_val:
            _set(ws, row, 1, font=LABEL_FONT, alignment=Alignment(horizontal="left", vertical="center", wrap_text=True))
        else:
            _set(ws, row, 1, alignment=Alignment(horizontal="left", vertical="center"))

        _set(ws, row, 2, font=DATA_FONT, alignment=Alignment(horizontal="left", vertical="center"))
        _set(
            ws,
            row,
            3,
            font=Font(name="Calibri", bold=True, size=11),
            fill=SUMMARY_FILL,
            alignment=Alignment(horizontal="center", vertical="center"),
            border=BORDER,
        )

        summary_peak = ws.cell(row, 3).value
        peak_target = int(summary_peak) if isinstance(summary_peak, (int, float)) else None

        for col in range(1, max_col + 1):
            cell = ws.cell(row, col)
            if col in weekend_cols and col >= 4:
                cell.fill = WEEKEND_FILL
            elif row_fill and col != 3:
                cell.fill = row_fill
            cell.border = BORDER

        for col in range(4, max_col + 1):
            val = ws.cell(row, col).value
            if not isinstance(val, (int, float)):
                continue
            iv = int(val)
            if iv == 0:
                _set(ws, row, col, font=ZERO_FONT, alignment=Alignment(horizontal="center"))
            elif peak_target and iv == peak_target:
                _set(
                    ws,
                    row,
                    col,
                    font=Font(name="Calibri", bold=True, size=11, color="006100"),
                    fill=PEAK_FILL,
                    alignment=Alignment(horizontal="center"),
                    border=Border(
                        left=Side(style="medium", color="006100"),
                        right=Side(style="medium", color="006100"),
                        top=Side(style="medium", color="006100"),
                        bottom=Side(style="medium", color="006100"),
                    ),
                )
            else:
                _set(ws, row, col, font=DATA_FONT, alignment=Alignment(horizontal="center"))

        ws.row_dimensions[row].height = 20

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 14
    for col in range(4, max_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 5.5


def _tenant_peak_hour_row(ws, detail_header_row: int, max_row: int, max_col: int) -> int | None:
  best_row = None
  best_count = -1
  for row in range(detail_header_row + 1, max_row + 1):
    vals = [ws.cell(row, c).value for c in range(1, 4)]
    if vals[0] in (None, "Date") or vals[1] in (None, "user_type"):
      continue
    count = ws.cell(row, 3).value
    if isinstance(count, (int, float)) and count > best_count:
      best_count = int(count)
      best_row = row
  return best_row


def _add_hourly_line_chart(ws, detail_header_row: int, data_row: int, max_col: int) -> None:
    if data_row is None or max_col < 27:
        return
    chart = LineChart()
    chart.title = "Hourly license profile (peak day)"
    chart.style = 10
    chart.y_axis.title = "Concurrent logins"
    chart.x_axis.title = "Hour of day"
    chart.height = 8
    chart.width = 16
    values = Reference(ws, min_col=4, min_row=data_row, max_col=max_col, max_row=data_row)
    chart.add_data(values, from_rows=True, titles_from_data=False)
    cats = Reference(ws, min_col=4, min_row=detail_header_row, max_col=max_col, max_row=detail_header_row)
    chart.set_categories(cats)
    chart.series[0].graphicalProperties.line.solidFill = "2E75B6"
    chart.series[0].graphicalProperties.line.width = 22000
    chart.series[0].marker.symbol = "circle"
    chart.series[0].marker.size = 5
    ws.add_chart(chart, f"A{data_row + 3}")




def _normalize_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _tenant_peak_dates(tenant_data: Any) -> set[date]:
    peak_dates: set[date] = set()
    if not tenant_data:
        return peak_dates
    for peak in tenant_data.peaks.values():
        if peak and getattr(peak, "peak_date", None):
            peak_dates.add(peak.peak_date)
    return peak_dates

def format_tenant_sheet(ws, tenant_data: Any = None, config: dict[str, Any] | None = None) -> None:
    max_col = ws.max_column
    max_row = ws.max_row
    if max_col < 3:
        return

    peak_header_row = None
    detail_header_row = None
    for row in range(1, min(max_row, 30) + 1):
        vals = [ws.cell(row, c).value for c in range(1, 4)]
        if vals == ["Date", "user_type", "Max_count"] and peak_header_row is None:
            peak_header_row = row
        if vals[:3] == ["Date", "user_type", "Max count"]:
            detail_header_row = row
            break

    if tenant_data and config:
        title = str(tenant_data.cfg.get("project_name", tenant_data.key)).replace("\n", " ")
        arc = tenant_data.cfg.get("arc", "")
        ws.insert_rows(1, 2)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(max_col, 10))
        _set(ws, 1, 1, value=title, font=TITLE_FONT, alignment=Alignment(horizontal="left", vertical="center"))
        _set(ws, 2, 1, value=f"{arc}  |  Peak-day hourly breakdown", font=SUBTITLE_FONT)
        if peak_header_row:
            peak_header_row += 2
        if detail_header_row:
            detail_header_row += 2
        max_row += 2

    if peak_header_row:
        for col in range(1, 4):
            _set(ws, peak_header_row, col, fill=HEADER_FILL, font=HEADER_FONT,
                 alignment=Alignment(horizontal="center", vertical="center"), border=BORDER)
        for row in range(peak_header_row + 1, (detail_header_row or peak_header_row + 4)):
            if _is_blank_row(ws, row, 3):
                continue
            for col in range(1, 4):
                _set(ws, row, col, fill=PEAK_FILL, font=Font(name="Calibri", bold=True, size=10, color="006100"),
                     border=BORDER, alignment=Alignment(horizontal="center", vertical="center"))

    peak_data_row = None
    if detail_header_row:
        for col in range(1, max_col + 1):
            _set(ws, detail_header_row, col, fill=HEADER_FILL, font=HEADER_FONT,
                 alignment=Alignment(horizontal="center", vertical="center", wrap_text=True), border=BORDER)
        ws.freeze_panes = ws.cell(detail_header_row + 1, 4).coordinate
        peak_dates = _tenant_peak_dates(tenant_data)
        agent_peak = tenant_data.peaks.get("agent") if tenant_data else None
        agent_peak_date = agent_peak.peak_date if agent_peak else None

        alt = False
        for row in range(detail_header_row + 1, max_row + 1):
            if _is_blank_row(ws, row, max_col):
                continue
            alt = not alt
            row_date = _normalize_date(ws.cell(row, 1).value)
            is_peak_day = bool(row_date and row_date in peak_dates)
            fill = PEAK_FILL if is_peak_day else (ALT_FILL if alt else None)
            row_font = Font(name="Calibri", bold=True, size=10, color="006100") if is_peak_day else DATA_FONT

            for col in range(1, max_col + 1):
                cell = ws.cell(row, col)
                if fill:
                    cell.fill = fill
                cell.font = row_font
                cell.border = BORDER
                cell.alignment = Alignment(horizontal="center", vertical="center")

            ws.row_dimensions[row].height = 18
            if is_peak_day:
                if agent_peak_date and row_date == agent_peak_date:
                    peak_data_row = row
                elif peak_data_row is None:
                    peak_data_row = row
                count_val = ws.cell(row, 3).value
                if isinstance(count_val, (int, float)):
                    if peak_data_row is None or int(count_val) > int(ws.cell(peak_data_row, 3).value or 0):
                        peak_data_row = row

        if peak_data_row:
            _add_hourly_line_chart(ws, detail_header_row, peak_data_row, max_col)

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 11
    for col in range(4, max_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 6.5


def _collect_dashboard_rows(tenants_data: list[Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for td in tenants_data:
        agent = td.peaks.get("agent")
        sup = td.peaks.get("supervisor")
        wb = td.peaks.get("wallboard")
        rows.append({
            "tenant": str(td.cfg.get("sheet_name", td.key)).replace("\n", " "),
            "arc": td.cfg.get("arc", ""),
            "agent_peak": agent.peak_count if agent else 0,
            "agent_date": agent.peak_date if agent and agent.peak_date else None,
            "sup_peak": sup.peak_count if sup else 0,
            "wb_peak": wb.peak_count if wb else 0,
            "daily_agent": dict(agent.daily_counts) if agent and agent.daily_counts else {},
        })
    rows.sort(key=lambda r: r["agent_peak"], reverse=True)
    return rows


def _daily_agent_totals(tenants_data: list[Any]) -> list[tuple[date, int]]:
    totals: dict[date, int] = defaultdict(int)
    for td in tenants_data:
        agent = td.peaks.get("agent")
        if not agent or not agent.daily_counts:
            continue
        for d, c in agent.daily_counts.items():
            totals[d] += int(c)
    return sorted(totals.items())


def add_dashboard_sheet(wb, tenants_data: list[Any], month: str, config: dict[str, Any]) -> None:
    if "Dashboard" in wb.sheetnames:
        del wb["Dashboard"]
    ws = wb.create_sheet("Dashboard", 0)

    month_text = _month_label(month)
    ws.merge_cells("A1:O1")
    _set(ws, 1, 1, value=f"TMONE License Utilization Dashboard", font=BANNER_FONT, fill=BANNER_FILL,
         alignment=Alignment(horizontal="center", vertical="center"))
    ws.row_dimensions[1].height = 36
    ws.merge_cells("A2:O2")
    _set(ws, 2, 1, value=f"{month_text}  |  Generated {datetime.now().strftime('%d-%b-%Y %H:%M')}",
         font=SUBTITLE_FONT, fill=BANNER_FILL, alignment=Alignment(horizontal="center"))

    dash_rows = _collect_dashboard_rows(tenants_data)
    total_agent_peak = sum(r["agent_peak"] for r in dash_rows)
    top = next((r for r in dash_rows if r["agent_peak"] > 0), None)
    arc1 = sum(1 for td in tenants_data if td.cfg.get("arc") == "ARC-1")
    arc2 = sum(1 for td in tenants_data if td.cfg.get("arc") == "ARC-2")

    kpis = [
        ("Tenants", len(tenants_data)),
        ("Combined agent peak", total_agent_peak),
        ("Highest peak", f"{top['agent_peak']} ({top['tenant']})" if top else "—"),
        ("ARC-1 / ARC-2", f"{arc1} / {arc2}"),
    ]
    col = 1
    for label, value in kpis:
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 2)
        ws.merge_cells(start_row=5, start_column=col, end_row=6, end_column=col + 2)
        _set(ws, 4, col, value=label, font=KPI_LABEL_FONT, fill=KPI_FILL, border=KPI_BORDER,
             alignment=Alignment(horizontal="center"))
        _set(ws, 5, col, value=value, font=KPI_VALUE_FONT, fill=KPI_FILL, border=BORDER,
             alignment=Alignment(horizontal="center", vertical="center", wrap_text=True))
        col += 3

    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 28

    table_row = 8
    headers = ["Tenant", "ARC", "Agent peak", "Peak date", "Supervisor peak"]
    for idx, hdr in enumerate(headers, start=1):
        _set(ws, table_row, idx, value=hdr, fill=HEADER_FILL, font=HEADER_FONT, border=BORDER,
             alignment=Alignment(horizontal="center"))
    for i, row in enumerate(dash_rows, start=table_row + 1):
        fill = ALT_FILL if i % 2 == 0 else None
        values = [row["tenant"], row["arc"], row["agent_peak"], row["agent_date"], row["sup_peak"]]
        for j, val in enumerate(values, start=1):
            _set(ws, i, j, value=val, font=DATA_FONT, border=BORDER,
                 alignment=Alignment(horizontal="center" if j > 1 else "left", vertical="center"))
            if fill:
                ws.cell(i, j).fill = fill

    if dash_rows:
        bar = BarChart()
        bar.type = "bar"
        bar.title = "Top tenants — agent peak licenses"
        bar.style = 10
        bar.y_axis.title = "Tenant"
        bar.x_axis.title = "Peak concurrent agents"
        bar.height = 12
        bar.width = 18
        chart_count = min(15, len([r for r in dash_rows if r["agent_peak"] > 0]) or len(dash_rows))
        data = Reference(ws, min_col=3, min_row=table_row, max_row=table_row + chart_count)
        cats = Reference(ws, min_col=1, min_row=table_row + 1, max_row=table_row + chart_count)
        bar.add_data(data, titles_from_data=True)
        bar.set_categories(cats)
        bar.shape = 4
        if bar.series:
            bar.series[0].graphicalProperties.solidFill = "2E75B6"
            bar.series[0].dLbls = DataLabelList()
            bar.series[0].dLbls.showVal = True
        ws.add_chart(bar, "G8")

    daily = _daily_agent_totals(tenants_data)
    if daily:
        start_col = 17
        _set(ws, table_row, start_col, value="Date", fill=HEADER_FILL, font=HEADER_FONT, border=BORDER)
        _set(ws, table_row, start_col + 1, value="Total agents", fill=HEADER_FILL, font=HEADER_FONT, border=BORDER)
        for i, (d, total) in enumerate(daily, start=table_row + 1):
            _set(ws, i, start_col, value=d, font=DATA_FONT, border=BORDER,
                 alignment=Alignment(horizontal="center"))
            _set(ws, i, start_col + 1, value=total, font=DATA_FONT, border=BORDER,
                 alignment=Alignment(horizontal="center"))

        line = LineChart()
        line.title = "Daily agent total (all tenants)"
        line.style = 12
        line.y_axis.title = "Licenses"
        line.x_axis.title = "Day"
        line.height = 12
        line.width = 18
        values = Reference(ws, min_col=start_col + 1, min_row=table_row, max_row=table_row + len(daily))
        cats = Reference(ws, min_col=start_col, min_row=table_row + 1, max_row=table_row + len(daily))
        line.add_data(values, titles_from_data=True)
        line.set_categories(cats)
        if line.series:
            line.series[0].graphicalProperties.line.solidFill = "1F4E79"
            line.series[0].graphicalProperties.line.width = 25000
            line.series[0].marker.symbol = "circle"
        ws.add_chart(line, "G24")

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    ws.sheet_view.showGridLines = False


def format_utilization_workbook(
    wb,
    month: str,
    tenants_data: list[Any] | None = None,
    config: dict[str, Any] | None = None,
) -> None:
    if tenants_data and config:
        add_dashboard_sheet(wb, tenants_data, month, config)
        name_map = {str(td.cfg.get("sheet_name", td.key))[:31]: td for td in tenants_data}
    else:
        name_map = {}

    if "Summary" in wb.sheetnames:
        format_summary_sheet(wb["Summary"], month)

    for name in wb.sheetnames:
        if name in ("Dashboard", "Summary"):
            continue
        format_tenant_sheet(wb[name], tenant_data=name_map.get(name), config=config)

def format_login_workbook(wb, month: str) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    data_font = Font(name="Calibri", size=10)
    alt_fill = PatternFill("solid", fgColor="F5F7FA")
    for name in wb.sheetnames:
        ws = wb[name]
        if ws.max_row < 1:
            continue
        for col in range(1, 5):
            cell = ws.cell(1, col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER
        ws.row_dimensions[1].height = 20
        ws.freeze_panes = "A2"
        for row in range(2, ws.max_row + 1):
            fill = alt_fill if row % 2 == 0 else None
            for col in range(1, 5):
                cell = ws.cell(row, col)
                if fill:
                    cell.fill = fill
                cell.font = data_font
                cell.border = BORDER
                cell.alignment = Alignment(horizontal="left" if col == 1 else "center", vertical="center")
        ws.column_dimensions["A"].width = 42
        ws.column_dimensions["B"].width = 24
        ws.column_dimensions["C"].width = 24
        ws.column_dimensions["D"].width = 16
